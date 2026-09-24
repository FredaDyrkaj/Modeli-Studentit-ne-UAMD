"""
Merge Db_Bach_FTI, Db_Master_FTI, Db_Bachelor_Buss, and Db_MASTER_Buss into one workbook.

First column (default: Niveli_studimit) labels each row's source. Column union: FTI Bachelor
header order, then new columns from FTI Master, then from Biznesi Bachelor, then from
Biznesi Master. Missing keys per row are written as blank.
"""

from __future__ import annotations

import argparse
import logging
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import Workbook

DEFAULT_BACH_FTI = Path("ResultFiles") / "Db_Bach_FTI.xlsx"
DEFAULT_MASTER_FTI = Path("ResultFiles") / "Db_Master_FTI.xlsx"
DEFAULT_BACH_BUSS = Path("ResultFiles") / "Db_Bachelor_Buss.xlsx"
DEFAULT_MASTER_BUSS = Path("ResultFiles") / "Db_MASTER_Buss.xlsx"
DEFAULT_OUTPUT = Path("ResultFiles") / "Db_Tot_Fti_Buss.xlsx"

SHEET_BACH_FTI = "Db_Bach_FTI"
SHEET_MASTER_FTI = "Db_Master_FTI"
SHEET_BACH_BUSS = "Db_Bachelor_Buss"
SHEET_MASTER_BUSS = "Db_MASTER_Buss"
SHEET_OUT = "Db_Tot_Fti_Buss"

# (label for first column, preferred sheet name in workbook)
SOURCES_4: Tuple[Tuple[str, str], ...] = (
    ("FTI Bachelor", SHEET_BACH_FTI),
    ("FTI Master", SHEET_MASTER_FTI),
    ("Biznesi Bachelor", SHEET_BACH_BUSS),
    ("Biznesi Master", SHEET_MASTER_BUSS),
)


def read_header_row(ws) -> List[Any]:
    return [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]


def union_headers_ordered(header_lists: List[List[Any]]) -> List[str]:
    seen: set = set()
    out: List[str] = []
    for hlist in header_lists:
        for h in hlist:
            if h is None or str(h).strip() == "":
                continue
            s = str(h).strip()
            if s in seen:
                continue
            seen.add(s)
            out.append(s)
    return out


def row_dict(ws, row: int, headers: List[str]) -> Dict[str, Any]:
    d: Dict[str, Any] = {}
    for ci, name in enumerate(headers, start=1):
        d[name] = ws.cell(row, ci).value
    return d


def load_sheet_rows(
    path: Path,
    sheet_name: str,
    log: logging.Logger,
) -> Tuple[List[str], List[Dict[str, Any]]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True, read_only=False)
    try:
        ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.active
        headers = [h for h in read_header_row(ws) if h is not None and str(h).strip() != ""]
        rows: List[Dict[str, Any]] = []
        for r in range(2, ws.max_row + 1):
            if all(ws.cell(r, c).value is None for c in range(1, len(headers) + 1)):
                continue
            rows.append(row_dict(ws, r, headers))
        return headers, rows
    finally:
        wb.close()


def merge_workbooks(
    sources: Tuple[Tuple[str, str], ...],
    paths: List[Path],
    out_path: Path,
    out_sheet: str,
    id_column: str,
    log: logging.Logger,
) -> None:
    if len(paths) != len(sources):
        raise ValueError("paths length must match sources")

    all_header_lists: List[List[Any]] = []
    all_rows: List[Tuple[str, List[Dict[str, Any]]]] = []

    for (label, sheet_name), path in zip(sources, paths):
        hdrs, rows = load_sheet_rows(path, sheet_name, log)
        all_header_lists.append(hdrs)
        all_rows.append((label, rows))
        log.info("Loaded %r: %d columns, %d rows.", path.name, len(hdrs), len(rows))

    merged_headers = union_headers_ordered(all_header_lists)
    full_headers = [id_column] + merged_headers
    total_rows = sum(len(r) for _, r in all_rows)
    log.info(
        "Union: %d columns (+ %s). Total data rows: %d.",
        len(merged_headers),
        id_column,
        total_rows,
    )

    out_wb = Workbook()
    ws_out = out_wb.active
    assert ws_out is not None
    ws_out.title = out_sheet
    ws_out.append(full_headers)

    for label, rows in all_rows:
        for rd in rows:
            line = [label]
            for h in merged_headers:
                line.append(rd.get(h))
            ws_out.append(line)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_wb.save(out_path)
    log.info("Wrote %s (%d data rows).", out_path, ws_out.max_row - 1)


def merge_all(
    paths: List[Path],
    out_path: Path,
    id_column: str,
    log: logging.Logger,
) -> None:
    merge_workbooks(SOURCES_4, paths, out_path, SHEET_OUT, id_column, log)


def main(argv: Optional[List[str]] = None) -> int:
    p = argparse.ArgumentParser(
        description="Merge FTI Bach/Master + Biznesi Bachelor/Master into Db_Tot_Fti_Buss.xlsx."
    )
    p.add_argument("--bach-fti", type=Path, default=DEFAULT_BACH_FTI, help="Db_Bach_FTI.xlsx")
    p.add_argument("--master-fti", type=Path, default=DEFAULT_MASTER_FTI, help="Db_Master_FTI.xlsx")
    p.add_argument("--bach-buss", type=Path, default=DEFAULT_BACH_BUSS, help="Db_Bachelor_Buss.xlsx")
    p.add_argument("--master-buss", type=Path, default=DEFAULT_MASTER_BUSS, help="Db_MASTER_Buss.xlsx")
    p.add_argument("--output", type=Path, default=DEFAULT_OUTPUT, help="Merged output xlsx")
    p.add_argument(
        "--id-column",
        default="Niveli_studimit",
        help="First column name (row source labels)",
    )
    p.add_argument("-v", "--verbose", action="store_true")
    args = p.parse_args(argv)

    logging.basicConfig(
        level=logging.DEBUG if args.verbose else logging.INFO,
        format="%(levelname)s %(message)s",
    )
    log = logging.getLogger("merge_db_tot_fti_buss")

    paths = [args.bach_fti, args.master_fti, args.bach_buss, args.master_buss]
    for path in paths:
        if not path.is_file():
            log.error("Input not found: %s", path)
            return 1

    merge_all(paths, args.output, args.id_column, log)
    return 0


# Backward compatibility: old name pointed at the 4-source tuple
SOURCES = SOURCES_4


if __name__ == "__main__":
    raise SystemExit(main())
