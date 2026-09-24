"""
Merge Db_Bach_FTI.xlsx and Db_Master_FTI.xlsx into one workbook.

Adds a first column (default: Niveli_studimit) with values Bachelor / Master.
Column union: Bachelor headers in file order, then any Master-only columns in Master order.
"""

from __future__ import annotations

import argparse
import logging
from pathlib import Path
from typing import Any, Dict, List, Optional

from openpyxl import Workbook

DEFAULT_BACH = Path(r"C:\\Users\\esoft\\Desktop\\Projekt-ModeliStudenteve\\FInalWork\\KodetPythonRregTeDhenash\\ResultFiles") / "Db_Bach_FTI.xlsx"
DEFAULT_MASTER = Path(r"C:\\Users\\esoft\\Desktop\\Projekt-ModeliStudenteve\\FInalWork\\KodetPythonRregTeDhenash\\ResultFiles") / "Db_Master_FTI.xlsx"
DEFAULT_OUTPUT = Path(r"C:\\Users\\esoft\\Desktop\\Projekt-ModeliStudenteve\\FInalWork\\KodetPythonRregTeDhenash\\ResultFiles") / "Db_Bach_Master_FTI.xlsx"

SHEET_BACH = "Db_Bach_FTI"
SHEET_MASTER = "Db_Master_FTI"
DEFAULT_SHEET_OUT = "Db_Bach_Master_FTI"


def read_header_row(ws) -> List[Any]:
    return [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]


def union_headers(bach_h: List[Any], master_h: List[Any]) -> List[str]:
    seen = set()
    out: List[str] = []
    for h in bach_h + master_h:
        if h is None or str(h).strip() == "":
            continue
        s = str(h).strip()
        if s in seen:
            continue
        seen.add(s)
        out.append(s)
    return out


def row_dict(ws, row: int, headers: List[str]) -> Dict[str, Any]:
    d: Dict[str, Any] = {}
    for ci, name in enumerate(headers, start=1):
        d[name] = ws.cell(row, ci).value
    return d


def merge_workbooks(
    bach_path: Path,
    master_path: Path,
    out_path: Path,
    id_column: str,
    log: logging.Logger,
) -> None:
    from openpyxl import load_workbook

    wb_b = load_workbook(bach_path, data_only=True, read_only=False)
    try:
        ws_b = wb_b[SHEET_BACH] if SHEET_BACH in wb_b.sheetnames else wb_b.active
        bach_headers = [h for h in read_header_row(ws_b) if h is not None and str(h).strip() != ""]
        bach_rows: List[Dict[str, Any]] = []
        for r in range(2, ws_b.max_row + 1):
            if all(ws_b.cell(r, c).value is None for c in range(1, len(bach_headers) + 1)):
                continue
            bach_rows.append(row_dict(ws_b, r, bach_headers))
    finally:
        wb_b.close()

    wb_m = load_workbook(master_path, data_only=True, read_only=False)
    try:
        ws_m = wb_m[SHEET_MASTER] if SHEET_MASTER in wb_m.sheetnames else wb_m.active
        master_headers = [h for h in read_header_row(ws_m) if h is not None and str(h).strip() != ""]
        master_rows: List[Dict[str, Any]] = []
        for r in range(2, ws_m.max_row + 1):
            if all(ws_m.cell(r, c).value is None for c in range(1, len(master_headers) + 1)):
                continue
            master_rows.append(row_dict(ws_m, r, master_headers))
    finally:
        wb_m.close()

    merged_headers = union_headers(bach_headers, master_headers)
    full_headers = [id_column] + merged_headers
    log.info(
        "Union: %d columns (+ %s). Bachelor rows: %d, Master rows: %d.",
        len(merged_headers),
        id_column,
        len(bach_rows),
        len(master_rows),
    )

    out_wb = Workbook()
    ws_out = out_wb.active
    assert ws_out is not None
    ws_out.title = DEFAULT_SHEET_OUT
    ws_out.append(full_headers)

    def emit(label: str, rows: List[Dict[str, Any]]) -> None:
        for rd in rows:
            line = [label]
            for h in merged_headers:
                line.append(rd.get(h))
            ws_out.append(line)

    emit("Bachelor", bach_rows)
    emit("Master", master_rows)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_wb.save(out_path)
    log.info("Wrote %s (%d data rows).", out_path, ws_out.max_row - 1)


def main(argv: Optional[List[str]] = None) -> int:
    p = argparse.ArgumentParser(
        description="Merge Db_Bach_FTI.xlsx and Db_Master_FTI.xlsx with a Bachelor/Master column."
    )
    p.add_argument("--bach", type=Path, default=DEFAULT_BACH, help="Path to Bachelor output xlsx")
    p.add_argument("--master", type=Path, default=DEFAULT_MASTER, help="Path to Master output xlsx")
    p.add_argument("--output", type=Path, default=DEFAULT_OUTPUT, help="Merged output xlsx path")
    p.add_argument(
        "--id-column",
        default="Niveli_studimit",
        help="Name of the first column (values: Bachelor, Master)",
    )
    p.add_argument("-v", "--verbose", action="store_true")
    args = p.parse_args(argv)

    logging.basicConfig(
        level=logging.DEBUG if args.verbose else logging.INFO,
        format="%(levelname)s %(message)s",
    )
    log = logging.getLogger("merge_db_fti_bach_master")

    if not args.bach.is_file():
        log.error("Bachelor file not found: %s", args.bach)
        return 1
    if not args.master.is_file():
        log.error("Master file not found: %s", args.master)
        return 1

    merge_workbooks(args.bach, args.master, args.output, args.id_column, log)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
